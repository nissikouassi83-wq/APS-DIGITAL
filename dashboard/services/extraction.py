"""Lecture d'un rapport hebdo SHA APS DIGITALE (.pptx) -> dictionnaires Python.

Ce module ne dépend pas de Django : on peut le tester seul
    python -m dashboard.services.extraction data/pptx/SHA_APS_DIGITALE_Sem_42.pptx
"""
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Variantes d'écriture d'un même motif -> libellé unique.
# Clé = libellé en minuscules, sans accents.
NORMALISATION = {
    "depot certificat de vie": "Dépôt certificat de vie",
    "depot de certificat de vie": "Dépôt certificat de vie",
    "modifications des coordonnees": "Modification des coordonnées",
    "demande de documents sur e-cnps": "Demande de document sur e-CNPS",
    "demande de document sur e-cnps": "Demande de document sur e-CNPS",
    "prestation cnps": "Prestations CNPS",
    "prestations cnps": "Prestations CNPS",
    "prestation d'accident de travail": "Prestations d'accident du travail",
    "prestations d'accident du travail": "Prestations d'accident du travail",
    "changement de situation famililale": "Changement de situation familiale",
    "standard": "Standard",
    "problemes d'immatriculation assure": "Problème d'immatriculation assuré",
    "probleme d'immatriculation": "Problème d'immatriculation assuré",
}


def _cle(texte):
    t = unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def normaliser_motif(motif):
    motif = re.sub(r"\s+", " ", motif).strip()
    return NORMALISATION.get(_cle(motif), motif)


def _nombre(pattern, texte):
    m = re.search(pattern, texte, re.I)
    return int(m.group(1)) if m else None


def _trouver_table(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_table:
                return shape.table
    return None


def _lire_tableau(table):
    motifs, globaux, dans_motifs = [], [], False
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        if cells[0] == "MOTIFS":
            dans_motifs = True
            continue
        try:
            valeurs = [int(x) for x in cells[1:5]]
        except ValueError:
            continue
        if cells[0].startswith("VOLUME GLOBAL"):
            globaux.append((cells[0], valeurs))
        elif dans_motifs and cells[0]:
            motifs.append((cells[0], valeurs))
    return motifs, globaux


def _zones_texte(slide):
    """Zones de texte libres d'un slide (hors zone « Réclamation client »)."""
    zones = []
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX and sh.text_frame.text.strip():
            if not sh.text_frame.text.startswith("Réclamation client"):
                zones.append(sh.text_frame.text)
    return zones


def extraire_rapport(chemin):
    chemin = Path(chemin)
    numero = int(re.search(r"Sem_(\d+)", chemin.name).group(1))
    prs = Presentation(str(chemin))

    # --- Centre de contact (slide 3)
    motifs_bruts, globaux = _lire_tableau(_trouver_table(prs))
    agreg = defaultdict(lambda: [0, 0, 0, 0])
    for libelle, v in motifs_bruts:
        a = agreg[normaliser_motif(libelle)]
        for i in range(4):
            a[i] += v[i]
    motifs = [dict(motif=m, appels=v[0], mails=v[1], messages=v[2], total=v[3])
              for m, v in agreg.items()]

    kpis = []
    for libelle, v in globaux:
        nom = libelle.replace("VOLUME GLOBAL DES ", "").strip()
        if nom == "VOLUME GLOBAL":
            nom = "TOTAL"
        kpis.append(dict(indicateur=nom, appels=v[0], mails=v[1], messages=v[2], total=v[3]))

    # --- Réclamations (slide 4)
    slide_recl = prs.slides[3]
    texte = ""
    for sh in slide_recl.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Réclamation client"):
            texte = sh.text_frame.text
    partie_semaine, _, partie_backlog = texte.partition("Backlogs")
    periode = re.search(r"du (\d\d/\d\d)\s*[–-]\s*(\d\d/\d\d/\d{4})", partie_semaine)
    reclamation = dict(
        periode=f"{periode.group(1)} - {periode.group(2)}" if periode else "",
        recues=_nombre(r"re[cç]ues\s*:\s*(\d+)", partie_semaine),
        recues_semaine_prec=_nombre(r"VS\s*(\d+)", partie_semaine),
        en_cours=_nombre(r"en cours de traitement\s*:\s*(\d+)", partie_semaine),
        resolues=_nombre(r"r[ée]solues\s*:\s*(\d+)", partie_semaine),
        backlog_recues=_nombre(r"re[cç]ues\s*:\s*(\d+)", partie_backlog),
        backlog_en_cours=_nombre(r"en cours de traitement\s*:\s*(\d+)", partie_backlog),
        backlog_resolues=_nombre(r"r[ée]solues\s*:\s*(\d+)", partie_backlog),
        hors_delai_25j=_nombre(r"25 jours\s*:\s*(\d+)", partie_backlog),
        taux_resolution_pct=_nombre(r"Taux de r[ée]solution\s*:\s*(\d+)", partie_backlog),
    )

    # --- Alertes et actions (1re zone = alertes, 2e = actions)
    alertes = []
    for idx, section in ((2, "Centre de contact"), (3, "Réclamations")):
        zones = _zones_texte(prs.slides[idx])
        for type_, zone in zip(("alerte", "action"), zones):
            for ligne in zone.split("\n"):
                if ligne.strip():
                    alertes.append(dict(section=section, type=type_, texte=ligne.strip()))

    return dict(numero=numero, periode=reclamation["periode"], fichier=chemin.name,
                motifs=motifs, kpis=kpis, reclamation=reclamation, alertes=alertes)


if __name__ == "__main__":
    r = extraire_rapport(sys.argv[1])
    print(f"Semaine {r['numero']} | {r['periode']} | {len(r['motifs'])} motifs | "
          f"{len(r['kpis'])} KPI | {len(r['alertes'])} textes")
    print(r["reclamation"])
